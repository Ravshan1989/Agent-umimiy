import datetime
import tempfile
import unittest
import zipfile
from pathlib import Path

from agent_umimiy import attachments, docs
from agent_umimiy.tools import Context, Toolbox

try:
    import openpyxl
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    HAVE = True
except ImportError:  # fixturalar uchun kutubxonalar faqat dasturchi muhitida kerak
    HAVE = False


@unittest.skipUnless(HAVE, "openpyxl/python-docx/python-pptx o'rnatilmagan")
class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def make_xlsx(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Savdo"
        ws.append(["Mijoz", "Narx", "Sana", "Vaqt", "To'langan", "Izoh"])
        ws.append(["Ali", 1200, datetime.date(2024, 3, 15), datetime.datetime(2024, 3, 15, 14, 30, 5), True, "yaxshi"])
        ws.append(["Vali", 3.5, datetime.date(2024, 4, 1), datetime.datetime(2024, 4, 1, 9, 0, 0), False, None])
        ws.append([None])  # bo'sh qator
        ws["D6"] = "siyrak"  # siyrak katak
        ws2 = wb.create_sheet("Ikkinchi")
        ws2.append(["a", "b"])
        ws2.append([1, 2])
        path = self.d / "t.xlsx"
        wb.save(path)
        return path

    def test_xlsx_values_dates_sparse_and_sheets(self):
        book = docs.Xlsx(self.make_xlsx())
        self.assertEqual(list(book.sheets), ["Savdo", "Ikkinchi"])
        rows = list(book.rows("Savdo"))
        self.assertEqual(rows[0], ["Mijoz", "Narx", "Sana", "Vaqt", "To'langan", "Izoh"])
        self.assertEqual(rows[1][:5], ["Ali", "1200", "2024-03-15", "2024-03-15 14:30:05", "TRUE"])
        self.assertEqual(rows[2][1], "3.5")
        self.assertEqual(rows[2][4], "FALSE")
        self.assertEqual(rows[3], [])  # 4- va 5-qatorlar bo'sh, lekin raqamlarni siljitmaydi
        self.assertEqual(rows[4], [])
        self.assertEqual(rows[5], ["", "", "", "siyrak"])  # D6: siyrak katak o'z ustunida
        with self.assertRaises(docs.DocError):
            list(book.rows("Yo'q"))

    def test_xlsx_text_and_analysis(self):
        text = docs.extract_text(self.make_xlsx())
        self.assertIn("=== Varaq: Savdo ===", text)
        self.assertIn("=== Varaq: Ikkinchi ===", text)
        self.assertIn("Ali\t1200\t2024-03-15", text)
        name, rows, sheets = docs.table_rows(self.d / "t.xlsx", "Ikkinchi")
        self.assertEqual((name, rows, sheets), ("Ikkinchi", [["a", "b"], ["1", "2"]], ["Savdo", "Ikkinchi"]))

    def test_docx_headings_lists_tables(self):
        doc = Document()
        doc.add_heading("Hisobot", 1)
        doc.add_heading("Natijalar", 2)
        doc.add_paragraph("Oddiy matn **emas**.")
        doc.add_paragraph("Birinchi band", style="List Bullet")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Oy", "Daromad", "Yanvar", "500"
        doc.save(self.d / "t.docx")
        text = docs.read_docx(self.d / "t.docx")
        self.assertIn("# Hisobot", text)
        self.assertIn("## Natijalar", text)
        self.assertIn("• Birinchi band", text)
        self.assertIn("| Oy | Daromad |", text)
        self.assertIn("| Yanvar | 500 |", text)

    def test_pptx_titles_body_tables_notes(self):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Reja"
        s.placeholders[1].text = "Birinchi nuqta\nIkkinchi nuqta"
        s.notes_slide.notes_text_frame.text = "Taqdimotda aytiladi"
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        s2.shapes.title.text = "Jadval"
        tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "A", "B", "1", "2"
        prs.save(self.d / "t.pptx")
        text = docs.read_pptx(self.d / "t.pptx")
        self.assertIn("--- Slayd 1: Reja ---", text)
        self.assertIn("Birinchi nuqta", text)
        self.assertIn("[Izoh: Taqdimotda aytiladi]", text)
        self.assertIn("--- Slayd 2: Jadval ---", text)
        self.assertIn("| A | B |", text)
        self.assertLess(text.index("Slayd 1"), text.index("Slayd 2"))

    def test_attachments_accept_office_and_reject_old(self):
        self.make_xlsx()
        self.assertIn("Ali", attachments.load(str(self.d / "t.xlsx")).text)
        for old in ("a.xls", "b.ppt", "c.doc"):
            (self.d / old).write_bytes(b"\xd0\xcf\x11\xe0 eski format")
            with self.assertRaises(attachments.AttachError) as cm:
                attachments.load(str(self.d / old))
            self.assertIn("eski", str(cm.exception))

    def test_broken_files_give_clean_errors(self):
        for name in ("x.xlsx", "x.docx", "x.pptx"):
            (self.d / name).write_bytes(b"bu zip emas")
            with self.assertRaises(docs.DocError):
                docs.extract_text(self.d / name)
        with zipfile.ZipFile(self.d / "bad.docx", "w") as z:
            z.writestr("word/document.xml", "<xml buzuq")
        with self.assertRaises(docs.DocError):
            docs.extract_text(self.d / "bad.docx")
        with zipfile.ZipFile(self.d / "bad.xlsx", "w") as z:
            z.writestr("boshqa.txt", "x")
        with self.assertRaises(docs.DocError):
            docs.extract_text(self.d / "bad.xlsx")


class TableTests(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def test_csv_delimiters_and_encodings(self):
        (self.d / "a.csv").write_text("ism;yosh\nAli;30\nVali;25\n", encoding="utf-8")
        self.assertEqual(docs.read_csv_rows(self.d / "a.csv"), [["ism", "yosh"], ["Ali", "30"], ["Vali", "25"]])
        (self.d / "b.csv").write_bytes("город,цена\nМосква,100\n".encode("cp1251"))
        self.assertEqual(docs.read_csv_rows(self.d / "b.csv")[1], ["Москва", "100"])
        (self.d / "c.tsv").write_text("x\ty\n1\t2\n")
        self.assertEqual(docs.read_csv_rows(self.d / "c.tsv")[1], ["1", "2"])

    def test_analyze_numbers_dates_text_empty(self):
        rows = [["shahar", "narx", "sana", "bo'sh"]] + [
            [("Toshkent" if i % 3 else "Samarqand"), str(100 + i), f"2024-01-{i + 1:02d}", ""] for i in range(10)]
        out = docs.analyze(rows)
        self.assertIn("Jadval: 10 qator × 4 ustun", out)
        self.assertIn("min 100", out)
        self.assertIn("maks 109", out)
        self.assertIn("o'rtacha 104.5", out)
        self.assertIn("sana: 2024-01-01 … 2024-01-10", out)
        self.assertIn("Toshkent (6)", out)
        self.assertIn("bo'sh ustun", out)

    def test_analyze_handles_odd_numbers(self):
        out = docs.analyze([["x"], ["1 200"], ["3,5"], ["1,234.5"], ["abc"]])
        self.assertIn("matn", out) if "son:" not in out else self.assertIn("son:", out)  # aralash ustun yiqilmaydi
        self.assertEqual(docs.analyze([]), "Jadval bo'sh.")

    def test_big_table_gets_summary_not_wall_of_text(self):
        rows = ["id,qiymat,izoh"] + [f"{i},{i * 2},qator raqami {i} uzun izoh matni shu yerda" for i in range(6000)]
        (self.d / "big.csv").write_text("\n".join(rows))
        text = docs.extract_text(self.d / "big.csv")
        self.assertIn("Jadval katta (6001 qator)", text)
        self.assertIn("Jadval: 6000 qator × 3 ustun", text)
        self.assertLess(len(text), 60_000)


class ToolTests(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        rows = ["id,shahar,summa"] + [f"{i},{'Toshkent' if i % 2 else 'Buxoro'},{i * 10}" for i in range(1, 201)]
        (self.d / "data.csv").write_text("\n".join(rows))
        self.tb = Toolbox(Context(root=str(self.d), skills=False), "t")

    def test_analyze_and_read_table(self):
        out, err = self.tb.run("analyze_data", {"path": "data.csv"})
        self.assertFalse(err, out)
        self.assertIn("200 qator × 3 ustun", out)
        self.assertIn("jami 201000", out)
        out, err = self.tb.run("read_table", {"path": "data.csv", "offset": 3, "limit": 2})
        self.assertIn("[Jami 201 qator; 3–4-qatorlar", out)
        self.assertIn("     3  2\tBuxoro\t20", out)
        out, _ = self.tb.run("read_table", {"path": "data.csv", "contains": "buxoro", "limit": 3})
        self.assertIn("100 ta qator", out)
        self.assertIn("     3  2\tBuxoro", out)  # asl qator raqami (sarlavha = 1)
        out, _ = self.tb.run("read_table", {"path": "data.csv", "contains": "yo'q-matn"})
        self.assertIn("0 ta qator", out)

    def test_table_tools_errors_and_escape(self):
        self.assertTrue(self.tb.run("analyze_data", {"path": "yoq.csv"})[1])
        self.assertTrue(self.tb.run("analyze_data", {"path": "../x.csv"})[1])
        (self.d / "a.png").write_bytes(b"\x89PNG")
        out, err = self.tb.run("read_table", {"path": "a.png"})
        self.assertTrue(err)
        self.assertIn("csv", out)

    def test_read_file_streams_big_text_in_parts(self):
        big = self.d / "big.log"
        big.write_text("\n".join(f"qator {i} " + "x" * 80 for i in range(1, 60001)))  # ~5 MB
        self.assertGreater(big.stat().st_size, 3_000_000)
        out, err = self.tb.run("read_file", {"path": "big.log", "offset": 40001, "limit": 3})
        self.assertFalse(err)
        self.assertIn("Katta fayl: 60000 qator; 40001–40003", out)
        self.assertIn("40002  qator 40002 ", out)
        self.assertIn("offset=40004", out)
        self.assertNotIn("qator 40004", out)
        out, _ = self.tb.run("read_file", {"path": "big.log"})
        self.assertIn("1–", out)
        self.assertLess(len(out), 220_000)

    @unittest.skipUnless(HAVE, "fixtura kutubxonalari yo'q")
    def test_read_file_office(self):
        doc = Document()
        doc.add_heading("Sarlavha", 1)
        doc.save(self.d / "t.docx")
        out, err = self.tb.run("read_file", {"path": "t.docx"})
        self.assertEqual((out.strip(), err), ("# Sarlavha", False))
        wb = openpyxl.Workbook()
        wb.active.append(["a", "b"])
        wb.save(self.d / "t.xlsx")
        self.assertIn("a\tb", self.tb.run("read_file", {"path": "t.xlsx"})[0])


if __name__ == "__main__":
    unittest.main()
